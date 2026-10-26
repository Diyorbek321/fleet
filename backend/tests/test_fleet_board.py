"""The fleet board: one row per truck, where it is, and its export and import.

It replaces the sheet a dispatcher fills in by phone every morning, so what
must hold is what that sheet is read for: every truck on it in the desk's own
order, the place the GPS last saw it, the trip each way that is still under
way — never a delivered one — and a copy in the dispatchers' group each
morning, once.
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import settings
from app.models.dispatch_groups import DispatchTelegramGroup
from app.models.enums import TripStage, TripStatus
from app.models.trips import Trip
from app.models.trucks import TruckLocation
from app.routers import telegram as telegram_router
from app.services import fleet_board_digest, trip_orders
from app.services.telegram import SendResult

GROUP_CHAT = {"id": -100555, "type": "supergroup", "title": "Dispetcherlar"}


async def _truck(client: AsyncClient, headers, plate: str) -> str:
    res = await client.post("/api/trucks", headers=headers, json={"name": plate, "plate_number": plate})
    assert res.status_code in (200, 201), res.text
    return res.json()["id"]


async def _trip(client: AsyncClient, headers, truck_id: str, **fields) -> str:
    res = await client.post("/api/trips", headers=headers, json={"truck_id": truck_id, **fields})
    assert res.status_code == 200, res.text
    return res.json()["id"]


async def _advance(client: AsyncClient, headers, trip_id: str, to: str) -> None:
    res = await client.post(f"/api/trips/{trip_id}/advance", headers=headers, json={"to_status": to})
    assert res.status_code == 200, res.text


async def _board(client: AsyncClient, headers) -> list[dict]:
    res = await client.get("/api/fleet-board", headers=headers)
    assert res.status_code == 200, res.text
    return res.json()["rows"]


async def test_every_truck_is_listed_in_the_desks_order(client: AsyncClient, admin_headers):
    first = await _truck(client, admin_headers, "10 011 ECA/10 0894 BA")
    second = await _truck(client, admin_headers, "10 033 RBA/10 2462 BA")
    res = await client.put(
        "/api/trucks/order", headers=admin_headers, json={"truck_ids": [second, first]}
    )
    assert res.status_code == 200, res.text

    rows = await _board(client, admin_headers)

    assert [r["plate"] for r in rows] == ["10 033 RBA/10 2462 BA", "10 011 ECA/10 0894 BA"]
    assert [r["position"] for r in rows] == [1, 2]
    assert rows[0]["export_trip"] is None and rows[0]["import_trip"] is None


async def test_the_place_is_the_city_the_gps_last_saw(client: AsyncClient, admin_headers, db):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    db.add(
        TruckLocation(
            truck_id=truck_id, latitude=42.3, longitude=69.6, speed=0,
            address="Казахстан, Шымкент",
            recorded_at=datetime(2026, 10, 14, 3, 0, tzinfo=timezone.utc),
        )
    )
    await db.commit()

    row = (await _board(client, admin_headers))[0]

    # The sheet's column holds a town, not "country, town".
    assert row["location"] == "Шымкент"
    assert row["location_at"].startswith("2026-10-14")


async def test_export_and_import_land_in_their_own_columns(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    await _trip(
        client, admin_headers, truck_id, direction="uz_ru",
        origin_name="Ташкент", destination_name="Мелеуз",
        scheduled_start="2026-10-06T00:00:00Z", scheduled_end="2026-10-12T18:59:59Z",
    )
    await _trip(
        client, admin_headers, truck_id, direction="ru_uz",
        origin_name="Тобольск", destination_name="Янгиюль", is_loaded=False,
    )

    row = (await _board(client, admin_headers))[0]

    assert row["export_trip"]["route"] == "Ташкент - Мелеуз"
    assert row["export_trip"]["loading_date"] == "2026-10-06"
    # 18:59:59Z is the panel's "end of the day" in Tashkent: still the 12th there.
    assert row["export_trip"]["unloading_date"] == "2026-10-12"
    assert row["import_trip"]["route"] == "Тобольск - Янгиюль"
    assert row["import_trip"]["is_loaded"] is False
    assert row["import_trip"]["unloading_date"] is None


async def test_a_delivered_or_cancelled_trip_leaves_the_board(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    done = await _trip(client, admin_headers, truck_id, direction="uz_ru", origin_name="A", destination_name="B")
    dropped = await _trip(client, admin_headers, truck_id, direction="ru_uz", origin_name="C", destination_name="D")
    await _advance(client, admin_headers, done, "delivered")
    await _advance(client, admin_headers, dropped, "cancelled")

    row = (await _board(client, admin_headers))[0]

    assert row["export_trip"] is None
    assert row["import_trip"] is None


async def test_the_trip_on_the_road_wins_over_the_next_one_booked(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    moving = await _trip(client, admin_headers, truck_id, direction="ru_uz", origin_name="Чел", destination_name="Ташкент")
    await _advance(client, admin_headers, moving, "en_route")
    # Booked after it, so newer — and still not the one the sheet is about.
    await _trip(client, admin_headers, truck_id, direction="ru_uz", origin_name="Омск", destination_name="Самарканд")

    row = (await _board(client, admin_headers))[0]

    assert row["import_trip"]["route"] == "Чел - Ташкент"


async def test_the_loading_date_is_the_day_it_was_loaded_once_known(
    client: AsyncClient, admin_headers, db
):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    trip_id = await _trip(
        client, admin_headers, truck_id, direction="ru_uz", origin_name="Екат", destination_name="Ташкент",
        scheduled_start="2026-10-10T00:00:00Z",
    )
    trip = (await db.execute(select(Trip).where(Trip.id == trip_id))).scalar_one()
    trip.loaded_at = datetime(2026, 10, 12, 8, 0, tzinfo=timezone.utc)
    await db.commit()

    row = (await _board(client, admin_headers))[0]

    assert row["import_trip"]["loading_date"] == "2026-10-12"


async def test_a_disabled_truck_is_left_off(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    res = await client.put(f"/api/trucks/{truck_id}", headers=admin_headers, json={"is_enabled": False})
    assert res.status_code == 200, res.text

    assert await _board(client, admin_headers) == []


async def test_the_board_is_a_desk_view_not_a_drivers(client: AsyncClient, driver_login):
    for path in ("/api/fleet-board", "/api/fleet-board.xlsx", "/api/fleet-board/telegram-group"):
        assert (await client.get(path, headers=driver_login["headers"])).status_code == 403


async def test_the_spreadsheet_matches_the_board(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "10 620 RCA/10 4096 BA")
    await _trip(client, admin_headers, truck_id, direction="uz_ru", origin_name="Ташкент", destination_name="Уфа",
                scheduled_start="2026-10-08T00:00:00Z")

    res = await client.get("/api/fleet-board.xlsx", headers=admin_headers)

    assert res.status_code == 200, res.text
    assert "attachment" in res.headers["content-disposition"]
    ws = load_workbook(BytesIO(res.content)).active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "ЭКСПОРТ" in cells and "ИМПОРТ" in cells
    assert "10 620 RCA/10 4096 BA" in cells
    assert "Ташкент - Уфа" in cells
    assert "08.10.2026" in cells
    # Excel's "06.01.1900" — a blank date read as day zero — is what the hand
    # sheet prints for a truck with no import. Here a blank stays blank.
    assert not any("1900" in c for c in cells)


# ── The dispatchers' group ───────────────────────────────────────────────


@pytest.fixture
def telegram(monkeypatch):
    monkeypatch.setattr(settings, "telegram_bot_token", "TEST:token", raising=False)
    monkeypatch.setattr(settings, "telegram_webhook_secret", "", raising=False)
    monkeypatch.setattr(settings, "telegram_bot_username", "FleetBot", raising=False)
    messages: list[tuple[str, str]] = []
    documents: list[tuple[str, str, bytes]] = []

    async def _fake_send(chat_id, text, *, disable_notification=False):
        messages.append((str(chat_id), text))
        return SendResult(ok=True, status_code=200)

    async def _fake_document(chat_id, filename, content, caption):
        documents.append((str(chat_id), filename, content))
        return SendResult(ok=True, status_code=200)

    monkeypatch.setattr(trip_orders, "send_message", _fake_send)
    monkeypatch.setattr(telegram_router, "send_message", _fake_send)
    monkeypatch.setattr(fleet_board_digest, "send_document", _fake_document)
    return {"messages": messages, "documents": documents}


async def _webhook(client: AsyncClient, update: dict) -> None:
    res = await client.post("/api/telegram/webhook", json={"update_id": 1, **update})
    assert res.status_code == 200, res.text


async def _link_group(client: AsyncClient, headers) -> None:
    link = (await client.post("/api/fleet-board/telegram-group", headers=headers)).json()
    assert link["status"] == "pending"
    assert link["deep_link"].startswith("https://t.me/FleetBot?startgroup=dispatch_")
    token = link["deep_link"].split("startgroup=dispatch_", 1)[1]
    await _webhook(
        client, {"message": {"message_id": 1, "chat": GROUP_CHAT, "text": f"/start@FleetBot dispatch_{token}"}}
    )


async def test_adding_the_bot_with_the_link_binds_the_dispatchers_group(
    client: AsyncClient, admin_headers, telegram
):
    await _link_group(client, admin_headers)

    status = (await client.get("/api/fleet-board/telegram-group", headers=admin_headers)).json()
    assert status["status"] == "linked"
    assert status["chat_title"] == "Dispetcherlar"
    assert status["deep_link"] is None
    assert telegram["messages"] and telegram["messages"][0][0] == str(GROUP_CHAT["id"])


def _tashkent(hour: int, day: int = 14) -> datetime:
    # 05:00 Tashkent is 00:00 UTC.
    return datetime(2026, 10, day, hour - 5, 0, tzinfo=timezone.utc)


async def test_the_morning_copy_goes_to_the_group_once_a_day(
    client: AsyncClient, admin_headers, telegram, db
):
    await _truck(client, admin_headers, "01A123BC")
    await _link_group(client, admin_headers)

    assert await fleet_board_digest.run(db, now=_tashkent(6)) == 0  # too early
    assert await fleet_board_digest.run(db, now=_tashkent(8)) == 1
    assert await fleet_board_digest.run(db, now=_tashkent(8)) == 0  # same morning
    assert await fleet_board_digest.run(db, now=_tashkent(9)) == 0  # still the same day

    chat_id, filename, content = telegram["documents"][0]
    assert chat_id == str(GROUP_CHAT["id"])
    assert filename.endswith(".xlsx") and "2026-10-14" in filename
    assert load_workbook(BytesIO(content)).active is not None

    assert await fleet_board_digest.run(db, now=_tashkent(8, day=15)) == 1
    assert len(telegram["documents"]) == 2


async def test_no_group_means_no_morning_copy(client: AsyncClient, admin_headers, telegram, db):
    await _truck(client, admin_headers, "01A123BC")
    assert await fleet_board_digest.run(db, now=_tashkent(8)) == 0
    assert telegram["documents"] == []


async def test_a_group_the_bot_was_removed_from_gets_nothing(
    client: AsyncClient, admin_headers, telegram, db
):
    await _link_group(client, admin_headers)
    await _webhook(client, {"my_chat_member": {"chat": GROUP_CHAT, "new_chat_member": {"status": "left"}}})

    status = (await client.get("/api/fleet-board/telegram-group", headers=admin_headers)).json()
    assert status["status"] == "lost"
    assert await fleet_board_digest.run(db, now=_tashkent(8)) == 0


async def test_the_send_today_button_sends_now(client: AsyncClient, admin_headers, telegram):
    await _truck(client, admin_headers, "01A123BC")
    await _link_group(client, admin_headers)

    res = await client.post("/api/fleet-board/telegram-group/send", headers=admin_headers)

    assert res.status_code == 200, res.text
    assert len(telegram["documents"]) == 1


def test_board_date_is_unaffected_by_utc_midnight():
    # 23:30 UTC on the 13th is already the 14th in Tashkent.
    assert fleet_board_digest.local_day(datetime(2026, 10, 13, 23, 30, tzinfo=timezone.utc)) == date(2026, 10, 14)


async def test_a_test_send_before_the_morning_does_not_cancel_it(
    client: AsyncClient, admin_headers, telegram, db
):
    await _link_group(client, admin_headers)
    row = (await db.execute(select(DispatchTelegramGroup))).scalar_one()

    assert await fleet_board_digest.send_board(db, row, now=_tashkent(7)) is True
    assert await fleet_board_digest.run(db, now=_tashkent(8)) == 1


async def test_a_spent_link_cannot_bind_another_group_once_the_bot_is_gone(
    client: AsyncClient, admin_headers, telegram
):
    link = (await client.post("/api/fleet-board/telegram-group", headers=admin_headers)).json()
    token = link["deep_link"].split("startgroup=dispatch_", 1)[1]
    await _webhook(client, {"message": {"message_id": 1, "chat": GROUP_CHAT, "text": f"/start dispatch_{token}"}})
    await _webhook(client, {"my_chat_member": {"chat": GROUP_CHAT, "new_chat_member": {"status": "kicked"}}})

    stranger = {"id": -100666, "type": "group", "title": "Not ours"}
    await _webhook(client, {"message": {"message_id": 2, "chat": stranger, "text": f"/start dispatch_{token}"}})

    status = (await client.get("/api/fleet-board/telegram-group", headers=admin_headers)).json()
    assert status["status"] == "lost"


async def test_a_rejected_file_does_not_unlink_the_group(
    client: AsyncClient, admin_headers, telegram, db, monkeypatch
):
    await _link_group(client, admin_headers)

    async def _bad_request(chat_id, filename, content, caption):
        return SendResult(ok=False, status_code=400, permanently_failed=True)

    monkeypatch.setattr(fleet_board_digest, "send_document", _bad_request)
    assert await fleet_board_digest.run(db, now=_tashkent(8)) == 0

    status = (await client.get("/api/fleet-board/telegram-group", headers=admin_headers)).json()
    assert status["status"] == "linked"


def test_typed_text_cannot_become_a_spreadsheet_formula():
    from app.services.fleet_board import BoardRow, BoardTrip
    from app.services.fleet_board_xlsx import build_workbook
    import uuid as _uuid

    trip = BoardTrip(
        trip_id=_uuid.uuid4(), reference="R", route='=HYPERLINK("http://x","y") - Уфа',
        loading_date=None, unloading_date=None, is_loaded=True, status="planned",
    )
    row = BoardRow(position=1, truck_id=_uuid.uuid4(), plate="=1+1", location="@x",
                   location_at=None, export_trip=trip, import_trip=None)

    ws = load_workbook(BytesIO(build_workbook([row], date(2026, 10, 14)))).active
    cells = {c.value: c.data_type for r in ws.iter_rows() for c in r if c.value is not None}
    assert cells["=1+1"] == "s"
    assert cells['=HYPERLINK("http://x","y") - Уфа'] == "s"


# ── The customer, and where a truck stands empty ─────────────────────────


async def test_each_trip_shows_its_customer(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    await _trip(client, admin_headers, truck_id, direction="uz_ru", origin_name="A",
                destination_name="B", customer="Artel Logistics")

    row = (await _board(client, admin_headers))[0]

    assert row["export_trip"]["customer"] == "Artel Logistics"


async def test_unloaded_in_russia_is_empty_for_import_in_that_city(
    client: AsyncClient, admin_headers
):
    """Unloaded the export in Russia: the truck now waits for an import load there."""
    truck_id = await _truck(client, admin_headers, "01A123BC")
    trip = await _trip(client, admin_headers, truck_id, direction="uz_ru",
                       origin_name="Ташкент", destination_name="Россия, Москва")
    await _advance(client, admin_headers, trip, "delivered")

    row = (await _board(client, admin_headers))[0]

    assert row["empty"]["kind"] == "import"
    assert row["empty"]["place"] == "Москва"
    assert row["empty"]["since"] is not None


async def test_unloaded_in_uzbekistan_is_empty_for_export(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    trip = await _trip(client, admin_headers, truck_id, direction="ru_uz",
                       origin_name="Тобольск", destination_name="Янгиюль")
    await _advance(client, admin_headers, trip, "delivered")

    row = (await _board(client, admin_headers))[0]

    assert row["empty"]["kind"] == "export"
    assert row["empty"]["place"] == "Янгиюль"


async def test_empty_from_the_moment_the_driver_reports_unloaded(
    client: AsyncClient, admin_headers, db
):
    """The trip stays open for papers after unloading, but the truck is free."""
    truck_id = await _truck(client, admin_headers, "01A123BC")
    trip_id = await _trip(client, admin_headers, truck_id, direction="uz_ru",
                          origin_name="Ташкент", destination_name="Уфа")
    trip = await db.get(Trip, trip_id)
    trip.current_stage = TripStage.unloaded
    trip.status = TripStatus.en_route
    await db.commit()

    row = (await _board(client, admin_headers))[0]

    assert row["export_trip"] is not None
    assert row["empty"] == {**row["empty"], "kind": "import", "place": "Уфа"}


async def test_a_booked_return_load_means_not_empty(client: AsyncClient, admin_headers):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    done = await _trip(client, admin_headers, truck_id, direction="uz_ru",
                       origin_name="Ташкент", destination_name="Москва")
    await _advance(client, admin_headers, done, "delivered")
    await _trip(client, admin_headers, truck_id, direction="ru_uz",
                origin_name="Москва", destination_name="Ташкент")

    row = (await _board(client, admin_headers))[0]

    assert row["empty"] is None


async def test_a_truck_with_no_trips_is_not_called_empty(client: AsyncClient, admin_headers):
    await _truck(client, admin_headers, "01A123BC")
    assert (await _board(client, admin_headers))[0]["empty"] is None


async def test_the_spreadsheet_names_the_customer_and_the_empty_city(
    client: AsyncClient, admin_headers
):
    truck_id = await _truck(client, admin_headers, "01A123BC")
    await _trip(client, admin_headers, truck_id, direction="uz_ru", origin_name="Ташкент",
                destination_name="Уфа", customer="Artel")
    other = await _truck(client, admin_headers, "01B456CD")
    trip = await _trip(client, admin_headers, other, direction="uz_ru",
                       origin_name="Ташкент", destination_name="Москва")
    await _advance(client, admin_headers, trip, "delivered")

    res = await client.get("/api/fleet-board.xlsx", headers=admin_headers)
    ws = load_workbook(BytesIO(res.content)).active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]

    assert "Заказчик" in cells
    assert "Artel" in cells
    assert any(c.startswith("Порожний — Москва") for c in cells)
