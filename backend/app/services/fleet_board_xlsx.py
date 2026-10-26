"""The fleet board as the spreadsheet dispatchers already keep.

Laid out like the hand sheet it replaces — a truck block, then ЭКСПОРТ and
ИМПОРТ side by side, each with route, loading date and expected unloading — in
the same colours, so the morning copy in the group reads at a glance by people
who have looked at that sheet every day. Russian, like the sheet.

One deliberate difference: a blank date stays blank. The hand sheet prints
"06.01.1900" wherever a formula met an empty cell.

A truck standing empty says so in the block it is waiting on — "Порожний —
Москва с 24.10" under ИМПОРТ — which is the sheet's "Пустые авто из РФ".
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.services.fleet_board import BoardEmpty, BoardRow, BoardTrip
from app.services.period_reports import report_tz

_HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
_TRUCK_FILL = PatternFill("solid", fgColor="A9D08E")
_PLACE_FILL = PatternFill("solid", fgColor="FFE699")
_EXPORT_FILL = PatternFill("solid", fgColor="C6E0B4")
_IMPORT_FILL = PatternFill("solid", fgColor="F8CBAD")
_DUE_FILL = PatternFill("solid", fgColor="C6EFCE")
_DUE_FONT = Font(bold=True, color="006100", size=9)
_BODY_FONT = Font(bold=True, size=9)
_HEAD_FONT = Font(bold=True, size=11)
_GROUP_FONT = Font(bold=True, size=13)
_THIN = Side(style="thin", color="808080")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center")

_COLUMNS = (
    ("№", 5),
    ("Данные ТС", 26),
    ("Место положения авто", 22),
    ("Заказчик", 20),
    ("Маршрут", 28),
    ("Дата погрузки", 12),
    ("Ожидаемая дата выгрузки", 14),
    ("Заказчик", 20),
    ("Маршрут", 28),
    ("Дата погрузки", 12),
    ("Ожидаемая дата выгрузки", 14),
)
_GROUPS = (("Транспорт", 1, 3), ("ЭКСПОРТ", 4, 7), ("ИМПОРТ", 8, 11))
# The expected-unloading columns, picked out as on the hand sheet.
_DUE_COLUMNS = (7, 11)
_TITLE_ROW, _GROUP_ROW, _HEAD_ROW = 1, 2, 3
_FORMULA_LEADS = ("=", "+", "-", "@")
_FIRST_DATA_ROW = 4


def _day(value: date | None) -> str | None:
    return value.strftime("%d.%m.%Y") if value else None


def _empty_text(empty: BoardEmpty) -> str:
    text = f"Порожний — {empty.place}" if empty.place else "Порожний"
    if empty.since:
        text += f" с {empty.since.astimezone(report_tz()).strftime('%d.%m')}"
    return text


def _trip_cells(
    trip: BoardTrip | None, empty: BoardEmpty | None
) -> tuple[str | None, str | None, str | None, str | None]:
    if trip is None:
        # Nothing booked this way; if the truck is waiting on this way, say so.
        return None, _empty_text(empty) if empty else None, None, None
    route = trip.route if trip.is_loaded else f"{trip.route} (порожний)"
    return trip.customer, route, _day(trip.loading_date), _day(trip.unloading_date)


def _header(ws, day: date) -> None:
    ws.cell(_TITLE_ROW, 1, f"Состояние автопарка на {_day(day)}").font = _GROUP_FONT
    for title, first, last in _GROUPS:
        ws.merge_cells(start_row=_GROUP_ROW, start_column=first, end_row=_GROUP_ROW, end_column=last)
        cell = ws.cell(_GROUP_ROW, first, title)
        cell.font, cell.alignment = _GROUP_FONT, _CENTER
        for col in range(first, last + 1):
            ws.cell(_GROUP_ROW, col).border = _BORDER
    for col, (label, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(_HEAD_ROW, col, label)
        cell.font, cell.fill, cell.alignment, cell.border = _HEAD_FONT, _HEAD_FILL, _CENTER, _BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[_HEAD_ROW].height = 45


def _waiting(row: BoardRow, kind: str) -> BoardEmpty | None:
    return row.empty if row.empty and row.empty.kind == kind else None


def _write_row(ws, row_no: int, row: BoardRow) -> None:
    values = (
        row.position,
        row.plate,
        row.location,
        *_trip_cells(row.export_trip, _waiting(row, "export")),
        *_trip_cells(row.import_trip, _waiting(row, "import")),
    )
    fills = (
        _TRUCK_FILL, _TRUCK_FILL, _PLACE_FILL,
        _EXPORT_FILL, _EXPORT_FILL, _EXPORT_FILL, _EXPORT_FILL,
        _IMPORT_FILL, _IMPORT_FILL, _IMPORT_FILL, _IMPORT_FILL,
    )
    for col, (value, fill) in enumerate(zip(values, fills), start=1):
        cell = ws.cell(row_no, col, value)
        # Plates and routes are typed by people, and openpyxl writes any "=…"
        # string as a formula. Stored as text, "=HYPERLINK(…)" stays harmless.
        if isinstance(value, str) and value.startswith(_FORMULA_LEADS):
            cell.data_type = "s"
        cell.fill, cell.font, cell.border = fill, _BODY_FONT, _BORDER
        cell.alignment = _LEFT if col in (2, 3) else _CENTER
    # The expected unloading is the date the desk plans the truck's next load
    # around, so it is the one picked out — as on the hand sheet.
    for col in _DUE_COLUMNS:
        cell = ws.cell(row_no, col)
        if cell.value:
            cell.fill, cell.font = _DUE_FILL, _DUE_FONT


def build_workbook(rows: list[BoardRow], day: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Автопарк"
    _header(ws, day)
    for offset, row in enumerate(rows):
        _write_row(ws, _FIRST_DATA_ROW + offset, row)
    last = _FIRST_DATA_ROW + max(len(rows), 1) - 1
    ws.auto_filter.ref = f"A{_HEAD_ROW}:K{last}"
    ws.freeze_panes = f"A{_FIRST_DATA_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename_for(day: date) -> str:
    return f"avtopark-{day.isoformat()}.xlsx"
